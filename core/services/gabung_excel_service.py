import os
import re
import pandas as pd


class GabungExcelService:

    UNITUP_MAP = {
        "14500": "TELUK SEGARA",
        "14510": "TAIS",
        "14600": "NUSA INDAH",
        "14610": "CURUP",
        "14620": "KEPAHIANG",
        "14630": "MUARA AMAN",
        "14640": "MANNA",
        "14660": "ARGA MAKMUR",
        "14680": "MUKO MUKO",
        "14690": "BINTUHAN",
    }

    @classmethod
    def process_files(cls, file_paths):
        all_data = []

        for file_path in file_paths:
            df = cls.process_file(file_path)

            if not df.empty:
                all_data.append(df)

        if not all_data:
            return pd.DataFrame()

        result_df = pd.concat(
            all_data,
            ignore_index=True
        )

        # Tambahkan nomor urut
        result_df.insert(
            0,
            "No",
            range(1, len(result_df) + 1)
        )

        return result_df

    @classmethod
    def process_file(cls, file_path):
        # Baca Excel berdasarkan ekstensi
        ext = os.path.splitext(file_path)[1].lower()

        if ext == ".xls":
            df = pd.read_excel(
                file_path,
                engine="xlrd",
                header=1
            )
        elif ext == ".xlsx":
            df = pd.read_excel(
                file_path,
                engine="openpyxl",
                header=1
            )
        else:
            raise ValueError(
                f"Format file tidak didukung: {ext}"
            )

        # Bersihkan nama kolom
        df.columns = [
            str(col).strip()
            for col in df.columns
        ]

        # Isi UNITUP yang kosong akibat merge cell di Excel
        if "UNITUP" in df.columns:
            df["UNITUP"] = df["UNITUP"].ffill()

            # Ganti kode UNITUP menjadi nama UNITUP
            df["UNITUP"] = (
                df["UNITUP"]
                .astype(str)
                .str.strip()
                .str.replace(r"\.0$", "", regex=True)
                .map(cls.UNITUP_MAP)
            )

        # Isi TANGGAL yang kosong akibat merge cell di Excel
        if "TANGGAL" in df.columns:
            df["TANGGAL"] = df["TANGGAL"].ffill()


        # Ambil Bulan dan Tahun dari tanggal
        bulan, tahun = cls.get_bulan_tahun(df)

        # Tambahkan kolom Bulan dan Tahun
        df.insert(0, "Bulan", bulan)
        df.insert(1, "Tahun", tahun)

        # Hapus baris yang seluruhnya kosong
        df = df.dropna(
            how="all"
        )

        # Hapus baris subtotal / Total
        # Baris total memiliki URAIAN kosong setelah merge cell dibaca oleh Pandas
        if "URAIAN" in df.columns:
            df = df.dropna(
                subset=["URAIAN"]
            )

        # Hapus Grand Total jika masih ada
        df = df[
            ~df.astype(str).apply(
                lambda row: row.str.contains(
                    "Grand Total",
                    case=False,
                    na=False
                ).any(),
                axis=1
            )
        ]

        return df.reset_index(drop=True)

    @classmethod
    def get_bulan_tahun(cls, df):
        """
        Mengambil Bulan dan Tahun berdasarkan kolom TANGGAL.
        """

        if "TANGGAL" not in df.columns:
            return "", ""

        tanggal = pd.to_datetime(
            df["TANGGAL"],
            errors="coerce",
            dayfirst=True
        )

        valid_dates = tanggal.dropna()

        if valid_dates.empty:
            return "", ""

        date_value = valid_dates.iloc[0]

        bulan = date_value.strftime("%B")
        tahun = date_value.year

        # Ubah nama bulan ke Bahasa Indonesia
        bulan_map = {
            "January": "Januari",
            "February": "Februari",
            "March": "Maret",
            "April": "April",
            "May": "Mei",
            "June": "Juni",
            "July": "Juli",
            "August": "Agustus",
            "September": "September",
            "October": "Oktober",
            "November": "November",
            "December": "Desember",
        }

        bulan = bulan_map.get(
            bulan,
            bulan
        )

        return bulan, tahun


    @classmethod
    def export_to_excel(cls, df, output_path):
        """
        Menyimpan hasil gabungan ke file Excel (.xlsx)
        dengan format yang rapi.
        """
        from openpyxl import load_workbook
        from openpyxl.styles import Font, Alignment
        from openpyxl.utils import get_column_letter

        # Simpan DataFrame ke Excel
        df.to_excel(
            output_path,
            index=False,
            sheet_name="Gabung Excel"
        )

        # Buka kembali untuk formatting
        wb = load_workbook(output_path)
        ws = wb["Gabung Excel"]

        # Format header
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )

        # Format seluruh data
        for row in ws.iter_rows():
            for cell in row:
                cell.alignment = Alignment(
                    vertical="center"
                )

        # Lebar kolom otomatis
        for column_cells in ws.columns:
            max_length = 0
            column_letter = get_column_letter(
                column_cells[0].column
            )

            for cell in column_cells:
                if cell.value is not None:
                    cell_length = len(str(cell.value))

                    if cell_length > max_length:
                        max_length = cell_length

            # Batas lebar agar tidak terlalu lebar
            adjusted_width = min(
                max_length + 2,
                40
            )

            ws.column_dimensions[
                column_letter
            ].width = adjusted_width

        # Lebar kolom No
        ws.column_dimensions["A"].width = 7

        # Tinggi header
        ws.row_dimensions[1].height = 24

        # Freeze header
        ws.freeze_panes = "A2"

        # Simpan
        wb.save(output_path)   