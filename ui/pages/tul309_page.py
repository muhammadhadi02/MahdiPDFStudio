"""
==========================================================
Mahdi PDF Studio
TUL 309 Page
==========================================================

Halaman pengolahan data TUL 309.

Fitur:
- Pemilihan banyak file Excel
- Menampilkan daftar file
- Menghapus file terpilih
- Membersihkan seluruh daftar
- Proses dan export rekap TUL 309

==========================================================
"""

import os

import customtkinter as ctk
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from tkinter import filedialog, messagebox

from themes.colors import Colors
from themes.fonts import Fonts
from themes.icons import Icons

from core.services.tul309_service import TUL309Service


class TUL309Page(ctk.CTkFrame):

    FILE_LIST_HEIGHT = 395

    def __init__(
        self,
        master,
        status_bar=None
    ):
        super().__init__(
            master,
            fg_color=Colors.WORKSPACE_BG,
        )

        self.status_bar = status_bar
        self.selected_files = []

        self.configure_layout()
        self.create_header()
        self.create_file_section()
        self.update_info_panel()

    # ==================================================
    # LAYOUT
    # ==================================================

    def configure_layout(self):

        self.grid_rowconfigure(
            0,
            weight=0
        )

        self.grid_rowconfigure(
            1,
            weight=1
        )

        self.grid_columnconfigure(
            0,
            weight=1
        )

    # ==================================================
    # HEADER
    # ==================================================

    def create_header(self):

        header_frame = ctk.CTkFrame(
            self,
            fg_color="transparent",
        )

        header_frame.grid(
            row=0,
            column=0,
            sticky="ew",
            padx=30,
            pady=(0, 0),
        )

        header_frame.grid_columnconfigure(
            0,
            weight=0
        )

        header_frame.grid_columnconfigure(
            1,
            weight=1
        )

        header_frame.grid_columnconfigure(
            2,
            weight=0
        )

        # ==================================================
        # SUMMARY CARD
        # ==================================================

        summary_card = ctk.CTkFrame(
            header_frame,
            width=350,
            height=110,
            corner_radius=16,
            fg_color="#FFFFFF",
            border_width=1,
            border_color=Colors.BORDER,
        )

        summary_card.grid(
            row=0,
            column=2,
            rowspan=2,
            sticky="ne",
            padx=(20, 0),
        )

        summary_card.grid_propagate(False)

        summary_title = ctk.CTkLabel(
            summary_card,
            text="▥  Ringkasan",
            font=ctk.CTkFont(
                size=16,
                weight="bold"
            ),
            text_color=Colors.TEXT_PRIMARY,
        )

        summary_title.pack(
            anchor="w",
            padx=22,
            pady=(5, 5),
        )

        # ==================================================
        # SUMMARY DATA
        # ==================================================

        summary_container = ctk.CTkFrame(
            summary_card,
            fg_color="transparent",
        )

        summary_container.pack(
            fill="x",
            padx=22,
            pady=(2, 2),
        )

        summary_container.grid_columnconfigure(
            0,
            weight=1
        )

        summary_container.grid_columnconfigure(
            1,
            weight=1
        )

        self.total_file_label = ctk.CTkLabel(
            summary_container,
            text="Total File : 0",
            font=Fonts.PAGE_SUBTITLE,
            text_color=Colors.TEXT_SECONDARY,
            anchor="w",
        )

        self.total_file_label.grid(
            row=0,
            column=0,
            sticky="w",
            padx=(0, 15),
        )

        self.total_data_label = ctk.CTkLabel(
            summary_container,
            text="Total Data : 0",
            font=Fonts.PAGE_SUBTITLE,
            text_color=Colors.TEXT_SECONDARY,
            anchor="e",
        )

        self.total_data_label.grid(
            row=0,
            column=1,
            sticky="e",
            padx=(15, 0),
        )

        # ==================================================
        # ICON
        # ==================================================

        icon_box = ctk.CTkFrame(
            header_frame,
            width=80,
            height=80,
            corner_radius=24,
            fg_color=Colors.MERGE_BG,
        )

        icon_box.grid(
            row=0,
            column=0,
            rowspan=2,
            padx=(0, 20),
        )

        icon_box.grid_propagate(False)

        icon_label = ctk.CTkLabel(
            icon_box,
            text="",
            image=Icons.TUL_LOGO,
        )

        icon_label.place(
            relx=0.5,
            rely=0.5,
            anchor="center",
        )

        # ==================================================
        # TITLE
        # ==================================================

        title = ctk.CTkLabel(
            header_frame,
            text="TUL 309",
            font=Fonts.PAGE_TITLE,
            text_color=Colors.TEXT_PRIMARY,
        )

        title.grid(
            row=0,
            column=1,
            sticky="sw",
            pady=(0, 2),
        )

        # ==================================================
        # SUBTITLE
        # ==================================================

        subtitle = ctk.CTkLabel(
            header_frame,
            text=(
                "Rekapitulasi data TUL 309 "
                "dari beberapa file Excel."
            ),
            font=Fonts.PAGE_SUBTITLE,
            text_color=Colors.TEXT_SECONDARY,
        )

        subtitle.grid(
            row=1,
            column=1,
            sticky="nw",
        )

    # ==================================================
    # FILE SECTION
    # ==================================================

    def create_file_section(self):

        # ==================================================
        # OUTER CARD
        # ==================================================

        section = ctk.CTkFrame(
            self,
            fg_color=Colors.WORKSPACE_BG,

        )

        section.grid(
            row=1,
            column=0,
            padx=30,
            pady=(10, 0),
            sticky="nsew",
        )

        section.grid_columnconfigure(
            0,
            weight=1
        )

        section.grid_rowconfigure(
            1,
            weight=1
        )

        # ==================================================
        # FILE LIST CONTAINER
        # ==================================================

        self.file_list_container = ctk.CTkFrame(
            section,
            fg_color="transparent",
        )

        self.file_list_container.grid(
            row=1,
            column=0,
            padx=25,
            pady=(10, 10),
            sticky="nsew",
        )

        self.file_list_container.grid_columnconfigure(
            0,
            weight=1
        )

        self.file_list_container.grid_rowconfigure(
            0,
            weight=1
        )

        # ==================================================
        # SCROLLABLE FILE LIST
        # ==================================================

        self.file_list = ctk.CTkScrollableFrame(
            self.file_list_container,
            fg_color=Colors.WORKSPACE_BG,
            corner_radius=14,
        )

        self.file_list.grid(
            row=0,
            column=0,
            sticky="nsew",
        )

        self.file_list.grid_columnconfigure(
            0,
            weight=1
        )

        # ==================================================
        # EMPTY STATE
        # ==================================================

        self.empty_state = ctk.CTkFrame(
            self.file_list_container,
            fg_color=Colors.MERGE_BG,
            corner_radius=14,
        )

        self.empty_state.place(
            relx=0.5,
            rely=0.5,
            relwidth=0.96,
            relheight=0.90,
            anchor="center",
        )

        # ==================================================
        # EMPTY CONTENT
        # ==================================================

        empty_content = ctk.CTkFrame(
            self.empty_state,
            fg_color="transparent",
            width=600,
            height=320,
        )

        empty_content.place(
            relx=0.5,
            rely=0.5,
            anchor="center",
        )
        empty_content.pack_propagate(False)
        
        # ==================================================
        # ICON
        # ==================================================

        empty_icon = ctk.CTkLabel(
            empty_content,
            text="",
            image=Icons.EXCEL_EMPTY,
        )

        empty_icon.pack(
            pady=(0, 10)
        )

        # ==================================================
        # TITLE
        # ==================================================

        empty_title = ctk.CTkLabel(
            empty_content,
            text="Belum ada file Excel",
            font=ctk.CTkFont(
                size=20,
                weight="bold"
            ),
            text_color=Colors.TEXT_PRIMARY,
        )

        empty_title.pack()

        # ==================================================
        # DESCRIPTION
        # ==================================================

        empty_subtitle = ctk.CTkLabel(
            empty_content,
            text=(
                "Belum ada file Excel TUL 309 yang ditambahkan.\n"
                "Klik tombol \"Tambah File\" untuk memilih file."
            ),
            font=Fonts.PAGE_SUBTITLE,
            text_color=Colors.TEXT_SECONDARY,
            justify="center",
        )

        empty_subtitle.pack(
            pady=(6, 18)
        )

    

        # ==================================================
        # BUTTON AREA
        # ==================================================

        button_frame = ctk.CTkFrame(
            section,
            fg_color="transparent",
        )

        button_frame.grid(
            row=2,
            column=0,
            sticky="ew",
            padx=25,
            pady=(0, 20),
        )

        button_frame.grid_columnconfigure(
            0,
            weight=0
        )

        button_frame.grid_columnconfigure(
            1,
            weight=0
        )

        button_frame.grid_columnconfigure(
            2,
            weight=1
        )

        button_frame.grid_columnconfigure(
            3,
            weight=0
        )

        # ==================================================
        # TAMBAH FILE
        # ==================================================

        self.add_button = ctk.CTkButton(
            button_frame,
            text="+  Tambah File",
            width=170,
            height=48,
            corner_radius=10,
            fg_color=Colors.PRIMARY,
            hover_color=Colors.PRIMARY_HOVER,
            font=ctk.CTkFont(
                size=14,
                weight="bold"
            ),
            command=self.add_files,
        )

        self.add_button.grid(
            row=0,
            column=0,
            sticky="w",
        )

        # ==================================================
        # HAPUS SEMUA
        # ==================================================

        self.clear_button = ctk.CTkButton(
            button_frame,
            text="Hapus Semua",
            width=170,
            height=48,
            corner_radius=10,
            fg_color="#FFFFFF",
            hover_color="#F5F8FC",
            border_width=1,
            border_color=Colors.BORDER,
            text_color=Colors.TEXT_PRIMARY,
            font=ctk.CTkFont(
                size=14
            ),
            command=self.clear_files,
        )

        self.clear_button.grid(
            row=0,
            column=1,
            padx=(15, 0),
            sticky="w",
        )

        # ==================================================
        # PROCESS
        # ==================================================

        self.process_button = ctk.CTkButton(
            button_frame,
            text="GABUNGKAN TUL 309",
            width=230,
            height=48,
            corner_radius=10,
            fg_color=Colors.TEXT_MUTED,
            hover_color=Colors.TEXT_MUTED,
            text_color="#FFFFFF",
            font=ctk.CTkFont(
                size=14,
                weight="bold"
            ),
            command=self.process_files,
        )

        self.process_button.grid(
            row=0,
            column=3,
            sticky="e",
        )

    # ==================================================
    # ADD FILES
    # ==================================================

    def add_files(self):

        files = filedialog.askopenfilenames(
            title="Pilih File Excel TUL 309",
            filetypes=[
                (
                    "Excel Files",
                    "*.xlsx *.xls"
                )
            ]
        )

        if not files:
            return

        added = 0

        for file_path in files:

            if file_path not in self.selected_files:

                self.selected_files.append(
                    file_path
                )

                added += 1

        self.refresh_file_list()

        if added:
            self.update_status(
                f"{added} file Excel berhasil ditambahkan"
            )

    # ==================================================
    # REFRESH FILE LIST
    # ==================================================

    def refresh_file_list(self):

        for widget in self.file_list.winfo_children():

            if widget != self.empty_state:
                widget.destroy()

        if not self.selected_files:

            self.empty_state.place(
                relx=0.5,
                rely=0.5,
                anchor="center",
                relwidth=0.96,
                relheight=0.90,
            )

        else:

            self.empty_state.place_forget()

            for index, file_path in enumerate(
                self.selected_files,
                start=1
            ):

                self.create_file_item(
                    index,
                    file_path
                )

        self.update_info_panel()

    # ==================================================
    # FILE ITEM
    # ==================================================

    def create_file_item(
        self,
        index,
        file_path
    ):

        item = ctk.CTkFrame(
            self.file_list,
            fg_color="#F8FAFC",
            corner_radius=10,
            border_width=1,
            border_color=Colors.BORDER,
        )

        item.pack(
            fill="x",
            padx=15,
            pady=6,
        )

        # ==================================================
        # COLUMN CONFIGURATION
        # ==================================================

        item.grid_columnconfigure(
            0,
            weight=0
        )

        item.grid_columnconfigure(
            1,
            weight=0
        )

        item.grid_columnconfigure(
            2,
            weight=1
        )

        item.grid_columnconfigure(
            3,
            weight=0
        )

        # ==================================================
        # NUMBER
        # ==================================================

        number = ctk.CTkLabel(
            item,
            text=str(index),
            width=35,
            font=ctk.CTkFont(
                size=13,
                weight="bold"
            ),
            text_color=Colors.TEXT_SECONDARY,
        )

        number.grid(
            row=0,
            column=0,
            padx=(12, 5),
            pady=12,
        )

        # ==================================================
        # EXCEL ICON
        # ==================================================

        icon = ctk.CTkLabel(
            item,
            text="",
            image=Icons.EXCEL,
        )

        icon.grid(
            row=0,
            column=1,
            sticky="w",
            padx=(5, 10),
            pady=12,
        )

        # ==================================================
        # FILE NAME
        # ==================================================

        filename = ctk.CTkLabel(
            item,
            text=os.path.basename(file_path),
            anchor="w",
            justify="left",
            font=ctk.CTkFont(
                size=13
            ),
            text_color=Colors.TEXT_PRIMARY,
        )

        filename.grid(
            row=0,
            column=2,
            sticky="w",
            padx=(5, 10),
            pady=12,
        )

        # ==================================================
        # REMOVE
        # ==================================================

        remove_button = ctk.CTkButton(
            item,
            text="×",
            width=35,
            height=30,
            corner_radius=8,
            fg_color="transparent",
            hover_color="#FEE2E2",
            text_color="#DC2626",
            font=ctk.CTkFont(
                size=18,
                weight="bold"
            ),
            command=lambda path=file_path:
                self.remove_file(path),
        )

        remove_button.grid(
            row=0,
            column=3,
            padx=(5, 10),
            pady=10,
        )

    # ==================================================
    # REMOVE FILE
    # ==================================================

    def remove_file(
        self,
        file_path
    ):

        if file_path in self.selected_files:

            self.selected_files.remove(
                file_path
            )

        self.refresh_file_list()

        self.update_status(
            f"{len(self.selected_files)} file Excel dipilih."
        )

    # ==================================================
    # CLEAR FILES
    # ==================================================

    def clear_files(self):

        if not self.selected_files:
            return

        confirm = messagebox.askyesno(
            "Hapus Semua",
            "Apakah Anda yakin ingin menghapus semua file?",
            parent=self
        )

        if not confirm:
            return

        self.selected_files.clear()

        self.refresh_file_list()

        self.update_status(
            "Daftar file TUL 309 dikosongkan"
        )

    # ==================================================
    # UPDATE INFO
    # ==================================================

    def update_info_panel(self):

        total_files = len(
            self.selected_files
        )

        self.total_file_label.configure(
            text=f"Total File : {total_files}"
        )

        self.total_data_label.configure(
            text="Total Data : 0"
        )

        if total_files > 0:

            self.process_button.configure(
                state="normal",
                fg_color=Colors.PRIMARY,
                hover_color=Colors.PRIMARY_HOVER,
            )

        else:

            self.process_button.configure(
                state="disabled",
                fg_color=Colors.TEXT_MUTED,
                hover_color=Colors.TEXT_MUTED,
            )

    # ==================================================
    # PROCESS
    # ==================================================

    def process_files(self):

        if not self.selected_files:

            messagebox.showwarning(
                "TUL 309",
                "Silakan pilih minimal satu file Excel.",
                parent=self
            )

            return

        confirm = messagebox.askyesno(
            "Proses TUL 309",
            (
                f"{len(self.selected_files)} file Excel "
                "akan diproses.\n\n"
                "Data akan direkap berdasarkan:\n"
                "S, R, P, T, C, dan L.\n\n"
                "Lanjutkan proses?"
            ),
            parent=self
        )

        if not confirm:
            return

        try:

            self.update_status(
                "Sedang memproses data TUL 309..."
            )

            result = TUL309Service.process_files(
                self.selected_files
            )

            if result.empty:

                messagebox.showwarning(
                    "TUL 309",
                    "Tidak ada data yang berhasil diproses.",
                    parent=self
                )

                return

            self.total_data_label.configure(
                text=f"Total Data : {len(result)}"
            )

            output_path = filedialog.asksaveasfilename(
                title="Simpan Rekap TUL 309",
                defaultextension=".xlsx",
                filetypes=[
                    (
                        "Excel Workbook",
                        "*.xlsx"
                    )
                ],
                initialfile="Rekap_TUL_309.xlsx"
            )

            if not output_path:

                self.update_status(
                    "Proses TUL 309 dibatalkan."
                )

                return

            # ==================================================
            # SIMPAN DATA
            # ==================================================

            result.to_excel(
                output_path,
                index=False,
                sheet_name="TUL 309"
            )

            # ==================================================
            # FORMATTING EXCEL
            # ==================================================

            workbook = load_workbook(
                output_path
            )

            worksheet = workbook["TUL 309"]

            # Header
            for cell in worksheet[1]:

                cell.font = Font(
                    bold=True
                )

                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center"
                )

            # Isi data
            for row in worksheet.iter_rows(
                min_row=2
            ):

                for cell in row:

                    cell.alignment = Alignment(
                        vertical="center"
                    )

            # Auto width
            for column_cells in worksheet.columns:

                max_length = 0

                column_letter = (
                    column_cells[0].column_letter
                )

                for cell in column_cells:

                    if cell.value is not None:

                        length = len(
                            str(cell.value)
                        )

                        if length > max_length:
                            max_length = length

                worksheet.column_dimensions[
                    column_letter
                ].width = max_length + 3

            # Lebar kolom No
            worksheet.column_dimensions[
                "A"
            ].width = 7

            # Tinggi header
            worksheet.row_dimensions[
                1
            ].height = 24

            # Freeze header
            worksheet.freeze_panes = "A2"

            workbook.save(
                output_path
            )

            # ==================================================
            # SUCCESS
            # ==================================================

            self.update_status(
                "Rekap TUL 309 berhasil dibuat."
            )

            messagebox.showinfo(
                "TUL 309 Berhasil",
                (
                    "Rekap TUL 309 berhasil dibuat.\n\n"
                    f"Jumlah file : {len(self.selected_files)}\n"
                    f"Jumlah baris : {len(result)}\n\n"
                    f"File disimpan di:\n{output_path}"
                ),
                parent=self
            )

        except Exception as error:

            self.update_status(
                "Terjadi kesalahan saat memproses TUL 309."
            )

            messagebox.showerror(
                "TUL 309 - Error",
                (
                    "Terjadi kesalahan saat memproses "
                    "file Excel.\n\n"
                    f"{error}"
                ),
                parent=self
            )

    # ==================================================
    # STATUS BAR
    # ==================================================

    def update_status(
        self,
        message
    ):

        if self.status_bar is None:
            return

        try:

            self.status_bar.set_status(
                message
            )

        except AttributeError:

            pass